#!/usr/bin/env python3
"""Создаёт templates/dogovor_template.docx — шаблон договора возмездного оказания услуг.

Самодостаточен: использует python-docx напрямую, не импортирует vanya.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

# Копия FIELDS из SPEC §2.5 (vanya/extract.py) — ключи плейсхолдеров берутся строго отсюда.
FIELDS = (
    "familiya", "imya", "otchestvo", "fio", "data_rozhdeniya", "pasport_seriya", "pasport_nomer",
    "pasport_vydan", "pasport_data", "pasport_kod", "adres_registracii", "inn", "ogrn", "kpp",
    "organizaciya", "yuridicheskiy_adres", "bank", "bik", "raschetny_schet", "korr_schet",
    "telefon", "email", "dolzhnost", "data_dogovora", "nomer_dogovora", "summa",
)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "templates" / "dogovor_template.docx"

FONT = "Times New Roman"
CENTER = WD_ALIGN_PARAGRAPH.CENTER
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY


def _set_font(run, size: int, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)


def para(doc: Document, text: str, size: int = 11, bold: bool = False,
         align=JUSTIFY, after: int = 6) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    _set_font(p.add_run(text), size, bold)


def cell_text(cell, text: str, size: int = 10, bold: bool = False) -> None:
    cell.text = ""  # очищает до одного пустого абзаца
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _set_font(p.add_run(line), size, bold)


def build() -> Document:
    doc = Document()

    # Формат A4, поля.
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(20)
    sec.left_margin = Mm(25)
    sec.right_margin = Mm(25)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)

    # Шапка.
    para(doc, "ДОГОВОР ВОЗМЕЗДНОГО ОКАЗАНИЯ УСЛУГ № {{nomer_dogovora}}",
         size=14, bold=True, align=CENTER, after=4)
    para(doc, "{{data_dogovora}}", size=12, align=CENTER, after=12)

    # Преамбула.
    para(doc,
         "{{organizaciya}}, именуемое в дальнейшем «Исполнитель», в лице {{dolzhnost}}, "
         "действующего на основании Устава, с одной стороны, и {{fio}}, именуемый в дальнейшем "
         "«Заказчик», с другой стороны, совместно именуемые «Стороны», заключили настоящий "
         "договор о нижеследующем:",
         after=12)

    # Раздел 1.
    para(doc, "1. ПРЕДМЕТ ДОГОВОРА", bold=True, after=4)
    para(doc, "1.1. Исполнитель обязуется по заданию Заказчика оказать услуги, а Заказчик "
              "обязуется принять и оплатить эти услуги в порядке и на условиях, предусмотренных "
              "настоящим договором.", after=8)

    # Раздел 2.
    para(doc, "2. ПРАВА И ОБЯЗАННОСТИ СТОРОН", bold=True, after=4)
    para(doc, "2.1. Исполнитель обязуется оказать услуги качественно и в согласованный Сторонами срок.")
    para(doc, "2.2. Заказчик обязуется предоставить Исполнителю необходимые документы и сведения, "
              "а также оплатить услуги в соответствии с настоящим договором.", after=8)

    # Раздел 3.
    para(doc, "3. СТОИМОСТЬ И ПОРЯДОК РАСЧЁТОВ", bold=True, after=4)
    para(doc, "3.1. Стоимость услуг составляет {{summa}} рублей.")
    para(doc, "3.2. Оплата производится Заказчиком в течение 5 (пяти) рабочих дней с момента "
              "подписания акта об оказании услуг путём перечисления денежных средств на расчётный "
              "счёт Исполнителя.", after=8)

    # Раздел 4.
    para(doc, "4. ОТВЕТСТВЕННОСТЬ СТОРОН", bold=True, after=4)
    para(doc, "4.1. За неисполнение или ненадлежащее исполнение обязательств по настоящему договору "
              "Стороны несут ответственность в соответствии с законодательством Российской Федерации.", after=8)

    # Раздел 5.
    para(doc, "5. СРОК ДЕЙСТВИЯ", bold=True, after=4)
    para(doc, "5.1. Настоящий договор вступает в силу с момента подписания и действует до полного "
              "исполнения Сторонами своих обязательств.", after=8)

    # Раздел 6.
    para(doc, "6. ПОРЯДОК РАЗРЕШЕНИЯ СПОРОВ", bold=True, after=4)
    para(doc, "6.1. Все споры и разногласия, возникающие из настоящего договора, Стороны разрешают "
              "путём переговоров. При недостижении согласия спор передаётся на рассмотрение суда "
              "в соответствии с законодательством Российской Федерации.", after=8)

    # Раздел 7 — реквизиты сторон (таблица).
    para(doc, "7. РЕКВИЗИТЫ СТОРОН", bold=True, after=4)
    table = doc.add_table(rows=10, cols=2)
    table.style = "Table Grid"

    cell_text(table.cell(0, 0), "Исполнитель", bold=True)
    cell_text(table.cell(0, 1), "Заказчик", bold=True)
    cell_text(table.cell(1, 0), "{{organizaciya}}")
    cell_text(table.cell(1, 1), "{{fio}}")
    cell_text(table.cell(2, 0), "ИНН {{inn}}, КПП {{kpp}}")
    cell_text(table.cell(2, 1), "{{familiya}} {{imya}} {{otchestvo}}")
    cell_text(table.cell(3, 0), "ОГРН {{ogrn}}")
    cell_text(table.cell(3, 1), "Дата рождения: {{data_rozhdeniya}}")
    cell_text(table.cell(4, 0), "{{yuridicheskiy_adres}}")
    cell_text(table.cell(4, 1), "Паспорт: {{pasport_seriya}} {{pasport_nomer}}")
    cell_text(table.cell(5, 0), "Банк: {{bank}}")
    cell_text(table.cell(5, 1), "Выдан: {{pasport_vydan}}")
    cell_text(table.cell(6, 0), "БИК {{bik}}")
    cell_text(table.cell(6, 1), "Дата выдачи: {{pasport_data}}")
    cell_text(table.cell(7, 0), "р/с {{raschetny_schet}}")
    cell_text(table.cell(7, 1), "Код подразделения: {{pasport_kod}}")
    cell_text(table.cell(8, 0), "к/с {{korr_schet}}")
    cell_text(table.cell(8, 1), "Адрес: {{adres_registracii}}")
    cell_text(table.cell(9, 0), "тел.: {{telefon}}, e-mail: {{email}}")
    cell_text(table.cell(9, 1), "тел.: {{telefon}}")

    # Подписи.
    para(doc, "", after=10)
    sig = doc.add_table(rows=2, cols=2)
    cell_text(sig.cell(0, 0), "Исполнитель", bold=True)
    cell_text(sig.cell(0, 1), "Заказчик", bold=True)
    cell_text(sig.cell(1, 0), "{{organizaciya}}\n______________ / ______________ /")
    cell_text(sig.cell(1, 1), "{{fio}}\n______________ / ______________ /")

    return doc


def collect_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main() -> int:
    doc = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    text = collect_text(doc)
    keys = sorted(set(re.findall(r"\{\{(\w+)\}\}", text)))
    total = len(re.findall(r"\{\{(\w+)\}\}", text))
    unknown = [k for k in keys if k not in FIELDS]

    print(f"Создан: {OUT}")
    print(f"Плейсхолдеров: {total} вхождений, {len(keys)} уникальных ключей")
    if unknown:
        print(f"ВНИМАНИЕ: ключи вне FIELDS: {unknown}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
