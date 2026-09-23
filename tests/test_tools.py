"""Тесты инструментов (vanya.tools) и детерминированного сценария fill_contract."""

from __future__ import annotations

from docx import Document

from vanya import scenarios, tools


def _ws(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    ws.mkdir()
    monkeypatch.setenv("VANYA_WORKSPACE", str(ws))
    return ws


def test_tool_schemas_shape():
    schemas = tools.tool_schemas()
    assert len(schemas) == 6
    names = {s["function"]["name"] for s in schemas}
    assert names == {
        "list_files",
        "read_document",
        "extract_requisites",
        "fill_contract",
        "find_risks",
        "write_report",
    }
    for s in schemas:
        assert s["type"] == "function"
        assert "description" in s["function"]
        assert "parameters" in s["function"]


def test_dispatch_unknown_tool():
    result = tools.dispatch("no_such_tool", {})
    assert result["ok"] is False
    assert "error" in result


def test_dispatch_missing_argument_is_error(tmp_path, monkeypatch):
    _ws(tmp_path, monkeypatch)
    result = tools.dispatch("read_document", {})
    assert result["ok"] is False
    assert "error" in result


def test_dispatch_read_document_ok(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    (ws / "doc.txt").write_text("содержимое документа", encoding="utf-8")
    result = tools.dispatch("read_document", {"filename": "doc.txt"})
    assert result["ok"] is True
    assert result["result"]["text"] == "содержимое документа"


def test_dispatch_read_document_error_path(tmp_path, monkeypatch):
    _ws(tmp_path, monkeypatch)
    result = tools.dispatch("read_document", {"filename": "nope.txt"})
    assert result["ok"] is False
    assert "error" in result


def test_list_files(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    (ws / "a.txt").write_text("x", encoding="utf-8")
    (ws / "b.docx").write_bytes(b"x")
    result = tools.dispatch("list_files", {})
    assert result["ok"] is True
    names = [f["name"] for f in result["result"]]
    assert "a.txt" in names
    assert "b.docx" in names
    assert result["result"][0]["ext"] in {"txt", "docx"}


def test_write_report(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    result = tools.dispatch("write_report", {"name": "report.txt", "content": "текст отчёта"})
    assert result["ok"] is True
    assert (ws / "report.txt").read_text(encoding="utf-8") == "текст отчёта"


def test_extract_requisites_tool(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    (ws / "rekv.txt").write_text("ИНН: 7701234567\nКПП: 770101001", encoding="utf-8")
    result = tools.dispatch("extract_requisites", {"filename": "rekv.txt"})
    assert result["ok"] is True
    assert result["result"]["fields"]["inn"] == "7701234567"
    assert "ogrn" in result["result"]["missing"]


def _make_template(path):
    doc = Document()
    doc.add_paragraph("Заказчик: {{organizaciya}}")
    doc.add_paragraph("ИНН: {{inn}}")
    doc.save(str(path))


def test_scenario_fill_contract(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    tpl_dir = tmp_path / "tpl"
    tpl_dir.mkdir()
    _make_template(tpl_dir / "dogovor_template.docx")
    monkeypatch.setattr(scenarios, "TEMPLATES_DIR", tpl_dir)

    (ws / "source.txt").write_text(
        "Наименование организации: ООО «Ромашка»\nИНН: 7701234567",
        encoding="utf-8",
    )

    result = scenarios.fill_contract("source.txt")
    assert result["ok"] is True
    assert result["fields"]["organizaciya"] == "ООО «Ромашка»"
    assert "inn" not in result["missing"]
    out = Document(result["out_path"])
    body = "\n".join(p.text for p in out.paragraphs)
    assert "ООО «Ромашка»" in body
    assert "7701234567" in body


def test_fill_contract_tool(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    tpl_dir = tmp_path / "tpl"
    tpl_dir.mkdir()
    _make_template(tpl_dir / "dogovor_template.docx")
    monkeypatch.setattr(scenarios, "TEMPLATES_DIR", tpl_dir)

    (ws / "source.txt").write_text(
        "Наименование организации: ООО «Вектор»\nИНН: 1234567890",
        encoding="utf-8",
    )

    result = tools.dispatch(
        "fill_contract", {"source_filename": "source.txt", "out_name": "dog.docx"}
    )
    assert result["ok"] is True
    assert result["result"]["ok"] is True
    assert (ws / "dog.docx").exists()


def test_fill_contract_missing_template(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    tpl_dir = tmp_path / "empty"
    tpl_dir.mkdir()
    monkeypatch.setattr(scenarios, "TEMPLATES_DIR", tpl_dir)
    (ws / "source.txt").write_text("ИНН: 7701234567", encoding="utf-8")

    result = scenarios.fill_contract("source.txt")
    assert result["ok"] is False
    assert "error" in result


def test_scenario_fill_contract_package_mode(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    tpl_dir = tmp_path / "tpl"
    tpl_dir.mkdir()
    _make_template(tpl_dir / "dogovor_template.docx")
    monkeypatch.setattr(scenarios, "TEMPLATES_DIR", tpl_dir)

    (ws / "a.txt").write_text("ИНН: 7701234567", encoding="utf-8")
    (ws / "b.txt").write_text("Наименование организации: ООО «Ромашка»", encoding="utf-8")

    result = scenarios.fill_contract()
    assert result["ok"] is True
    assert result["package_mode"] is True
    assert set(result["sources"]) == {"a.txt", "b.txt"}
    assert result["fields"]["inn"] == "7701234567"
    assert result["fields"]["organizaciya"] == "ООО «Ромашка»"


def test_fill_contract_tool_package_mode(tmp_path, monkeypatch):
    ws = _ws(tmp_path, monkeypatch)
    tpl_dir = tmp_path / "tpl"
    tpl_dir.mkdir()
    _make_template(tpl_dir / "dogovor_template.docx")
    monkeypatch.setattr(scenarios, "TEMPLATES_DIR", tpl_dir)

    (ws / "a.txt").write_text("ИНН: 7701234567", encoding="utf-8")

    result = tools.dispatch("fill_contract", {})
    assert result["ok"] is True
    assert result["result"]["ok"] is True
    assert result["result"]["package_mode"] is True
    assert result["result"]["sources"] == ["a.txt"]
