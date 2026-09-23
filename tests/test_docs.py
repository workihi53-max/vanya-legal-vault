"""Тесты работы с документами (vanya.docs)."""

from __future__ import annotations

import pytest
from docx import Document

from vanya import docs


def _make_pdf(path, text: str) -> None:
    """Собирает минимальный валидный PDF с одним текстовым объектом."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        (
            b"<< /Length "
            + str(len(content)).encode()
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF"
    )
    path.write_bytes(out)


def test_read_text_file(tmp_path):
    p = tmp_path / "doc.txt"
    p.write_text("Привет, мир", encoding="utf-8")
    assert docs.read_document(p) == "Привет, мир"


def test_read_markdown(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text("# Заголовок\nтекст", encoding="utf-8")
    assert "# Заголовок" in docs.read_document(p)


def test_read_docx(tmp_path):
    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("Параграф один")
    doc.add_paragraph("Параграф два")
    doc.save(str(p))
    text = docs.read_document(p)
    assert "Параграф один" in text
    assert "Параграф два" in text


def test_read_pdf(tmp_path):
    p = tmp_path / "doc.pdf"
    _make_pdf(p, "Hello PDF")
    text = docs.read_document(p)
    assert "Hello PDF" in text


def test_read_unsupported_raises(tmp_path):
    p = tmp_path / "doc.bin"
    p.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError):
        docs.read_document(p)


def test_read_folder(tmp_path):
    (tmp_path / "a.txt").write_text("текст A", encoding="utf-8")
    (tmp_path / "b.md").write_text("текст B", encoding="utf-8")
    (tmp_path / "c.bin").write_bytes(b"\x00\x01")
    result = docs.read_folder(tmp_path)
    assert "=== a.txt ===" in result
    assert "текст A" in result
    assert "=== b.md ===" in result
    assert "c.bin" not in result


def test_read_folder_missing_raises(tmp_path):
    with pytest.raises(ValueError):
        docs.read_folder(tmp_path / "nope")


def test_write_text(tmp_path):
    p = docs.write_text(tmp_path / "sub" / "out.txt", "содержимое")
    assert p.exists()
    assert p.read_text(encoding="utf-8") == "содержимое"


def _make_template_docx(path):
    doc = Document()
    # плейсхолдер, разорванный по двум runs
    para = doc.add_paragraph()
    para.add_run("{{familiya")
    para.add_run("}} — далее текст")
    # плейсхолдер в таблице
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "ИНН: {{inn}}"
    # плейсхолдер в колонтитуле
    doc.sections[0].header.paragraphs[0].text = "Дата: {{data_dogovora}}"
    # несовпадающий плейсхолдер
    doc.add_paragraph("Лишнее: {{neizvestnoe_pole}}")
    doc.save(str(path))
    return doc


def test_fill_docx_split_runs_and_tables(tmp_path):
    tpl = tmp_path / "tpl.docx"
    _make_template_docx(tpl)
    out = tmp_path / "out.docx"
    values = {"familiya": "Иванов", "inn": "7701234567", "data_dogovora": "01.06.2024"}
    result = docs.fill_docx(tpl, values, out)
    assert result == out
    assert out.exists()

    filled = Document(str(out))
    # разорванный плейсхолдер собран и заменён
    assert "Иванов" in filled.paragraphs[0].text
    assert "{{" not in filled.paragraphs[0].text
    # таблица
    assert "7701234567" in filled.tables[0].cell(0, 0).text
    # колонтитул
    assert "01.06.2024" in filled.sections[0].header.paragraphs[0].text
    # несовпадающий плейсхолдер -> пустая строка
    assert "{{neizvestnoe_pole}}" not in filled.paragraphs[-1].text


def test_docx_placeholders(tmp_path):
    tpl = tmp_path / "tpl.docx"
    _make_template_docx(tpl)
    names = docs.docx_placeholders(tpl)
    assert "familiya" in names
    assert "inn" in names
    assert "data_dogovora" in names
    assert "neizvestnoe_pole" in names
