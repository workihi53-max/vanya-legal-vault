"""Чтение и запись документов: txt/md/csv/json, docx, pdf."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

TEXT_EXT = {".txt", ".md", ".csv", ".json"}
_DOCX_PDF_EXT = {".docx", ".pdf"}
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n".join(parts)


def read_document(path: str | Path) -> str:
    """Читает .txt/.md/.csv/.json/.docx/.pdf; иначе — ValueError."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in TEXT_EXT:
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        return _read_docx(p)
    if ext == ".pdf":
        return _read_pdf(p)
    raise ValueError(f"Неподдерживаемый формат файла: {ext or 'без расширения'}")


def read_folder(folder: str | Path, limit_chars: int = 20000) -> str:
    """Конкатенация текстов файлов папки: '=== <имя> ===\\n<текст>'."""
    f = Path(folder)
    if not f.is_dir():
        raise ValueError(f"Папка не найдена: {f}")
    chunks: list[str] = []
    total = 0
    for item in sorted(f.iterdir()):
        if not item.is_file():
            continue
        if item.suffix.lower() not in (TEXT_EXT | _DOCX_PDF_EXT):
            continue
        try:
            text = read_document(item)
        except Exception:
            continue  # битые/нетекстовые — пропускаем
        block = f"=== {item.name} ===\n{text}"
        if total + len(block) > limit_chars:
            remaining = limit_chars - total
            if remaining > 0:
                chunks.append(block[:remaining])
            break
        chunks.append(block)
        total += len(block)
    return "\n\n".join(chunks)


def write_text(path: str | Path, content: str) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return p


def _replace_placeholders(text: str, values: dict[str, str]) -> str:
    def repl(m: re.Match) -> str:
        key = m.group(1).strip()
        return values.get(key, "")

    return _PLACEHOLDER_RE.sub(repl, text)


def _fill_paragraphs(paragraphs, values: dict[str, str]) -> None:
    for para in paragraphs:
        full = "".join(run.text for run in para.runs)
        if "{{" not in full:
            continue
        new_text = _replace_placeholders(full, values)
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""


def _fill_tables(tables, values: dict[str, str]) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                _fill_paragraphs(cell.paragraphs, values)


def fill_docx(template: str | Path, values: dict[str, str], out_path: str | Path) -> Path:
    """Заменяет {{key}} (в т.ч. разорванные по runs), в таблицах и колонтитулах."""
    tpl = Path(template)
    out = Path(out_path)
    doc = Document(str(tpl))
    _fill_paragraphs(doc.paragraphs, values)
    _fill_tables(doc.tables, values)
    for section in doc.sections:
        for container in (section.header, section.footer):
            _fill_paragraphs(container.paragraphs, values)
            _fill_tables(container.tables, values)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _collect_placeholders(paragraphs, found: list[str], seen: set[str]) -> None:
    for para in paragraphs:
        full = "".join(run.text for run in para.runs)
        for m in _PLACEHOLDER_RE.finditer(full):
            key = m.group(1).strip()
            if key not in seen:
                seen.add(key)
                found.append(key)


def docx_placeholders(template: str | Path) -> list[str]:
    """Возвращает список уникальных имён плейсхолдеров в порядке появления."""
    doc = Document(str(template))
    found: list[str] = []
    seen: set[str] = set()
    _collect_placeholders(doc.paragraphs, found, seen)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _collect_placeholders(cell.paragraphs, found, seen)
    for section in doc.sections:
        for container in (section.header, section.footer):
            _collect_placeholders(container.paragraphs, found, seen)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        _collect_placeholders(cell.paragraphs, found, seen)
    return found
